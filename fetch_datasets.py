# -*- coding: utf-8 -*-
"""批量抓取 + 预处理「不同性质」的开源数据集，统一成 npy 格式，给 D3 翻案备用。

可达性实测（2026-10-17）：
    archive.ics.uci.edu          OK
    files.grouplens.org          OK
    codeload.github.com          OK（tarball）
    huggingface.co               ✗ 连接失败
    raw.githubusercontent.com    ✗ 连接失败
所以走 UCI + GroupLens。HF 上有更好的文本集，但这台机器现在够不着。

覆盖的数据性质
    text            sms         短信 spam/ham（真实文本，词袋哈希）
    timeseries      power       家庭用电（60 分钟窗口 -> 负荷分位类）
    recommendation  ml100k      MovieLens 100k（用户口味 + 电影类型 -> 评分档）
    tabular         adult       人口普查收入（数值 + 类别混合）
    tabular         credit      信用卡违约
    tabular         cancer      乳腺癌（小样本基线）
"""
import io, json, os, re, sys, zipfile
import numpy as np
import requests

ROOT = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(ROOT, "data")
RAW = os.path.join(DATA, "raw")
os.makedirs(RAW, exist_ok=True)
H = {"User-Agent": "Mozilla/5.0"}

SRC = {
    "adult":  "https://archive.ics.uci.edu/static/public/2/adult.zip",
    "sms":    "https://archive.ics.uci.edu/static/public/228/sms+spam+collection.zip",
    "power":  "https://archive.ics.uci.edu/static/public/235/individual+household+electric+power+consumption.zip",
    "ml100k": "https://files.grouplens.org/datasets/movielens/ml-100k.zip",
    "credit": "https://archive.ics.uci.edu/static/public/350/default+of+credit+card+clients.zip",
    "cancer": "https://archive.ics.uci.edu/static/public/17/breast+cancer+wisconsin+diagnostic.zip",
}


def fetch(name):
    zp = os.path.join(RAW, name + ".zip")
    if not os.path.exists(zp):
        print("下载 %-8s ..." % name); sys.stdout.flush()
        r = requests.get(SRC[name], headers=H, timeout=600)
        r.raise_for_status()
        open(zp, "wb").write(r.content)
    return zipfile.ZipFile(zp)


def read_zip(z, want):
    for n in z.namelist():
        if n.lower().endswith(want):
            return z.read(n)
    raise KeyError(want)


def std(X):
    X = np.asarray(X, dtype=np.float64)
    return ((X - X.mean(0, keepdims=True)) / (X.std(0, keepdims=True) + 1e-9)).astype(np.float32)


def balance(X, y, per=None, seed=0, max_classes=None):
    rng = np.random.default_rng(seed)
    cls = sorted(set(y.tolist()))
    if max_classes and len(cls) > max_classes:
        cnt = {c: int((y == c).sum()) for c in cls}
        cls = sorted(cls, key=lambda c: -cnt[c])[:max_classes]
        m = np.isin(y, cls)
        X, y = X[m], y[m]
        remap = {c: i for i, c in enumerate(sorted(cls))}
        y = np.array([remap[int(t)] for t in y])
    if per is None:
        per = min(int((y == c).sum()) for c in sorted(set(y.tolist())))
    idx = []
    for c in sorted(set(y.tolist())):
        ci = np.flatnonzero(y == c)
        idx.append(rng.choice(ci, size=min(per, len(ci)), replace=False))
    idx = np.concatenate(idx)
    rng.shuffle(idx)
    return X[idx], y[idx]


RESULTS = {}


def prep_adult():
    z = fetch("adult")
    txt = None
    for n in ("adult.data", "adult.data.txt"):
        try:
            txt = read_zip(z, n).decode("latin-1"); break
        except KeyError:
            pass
    if txt is None:
        cand = [n for n in z.namelist() if "adult.data" in n]
        txt = z.read(cand[0]).decode("latin-1")
    rows = [l.split(",") for l in txt.splitlines() if l.strip() and "?" not in l]
    cont = [0, 2, 4, 10, 11]
    cats = [1, 3, 5, 6, 7, 8, 9, 12, 13]
    Xc = np.array([[float(r[i]) for i in cont] for r in rows])
    blocks = [Xc]
    for ci in cats:
        vals = sorted(set(r[ci].strip() for r in rows))
        m = {v: i for i, v in enumerate(vals)}
        blocks.append(np.array([[1.0 if m[r[ci].strip()] == j else 0.0 for j in range(len(vals))] for r in rows]))
    X = np.hstack(blocks)
    y = np.array([1 if r[-1].strip() == ">50K" else 0 for r in rows], dtype=np.int64)
    return "tabular", balance(std(X), y, per=6000, max_classes=2)


def prep_sms():
    z = fetch("sms")
    cand = [n for n in z.namelist() if n.lower().endswith("spam.csv") or "SMSSpamCollection" in n]
    raw = z.read(cand[0]).decode("latin-1")
    ys, texts = [], []
    for line in raw.splitlines():
        if not line.strip():
            continue
        parts = line.split("\t")
        if len(parts) < 2:
            continue
        lab, msg = parts[0].strip().lower(), parts[1]
        if lab not in ("spam", "ham"):
            continue
        ys.append(1 if lab == "spam" else 0); texts.append(msg)
    D = 256
    X = np.zeros((len(texts), D), dtype=np.float64)
    for i, t in enumerate(texts):
        for w in re.findall(r"[a-z0-9']+", t.lower()):
            X[i, hash(w) % D] += 1.0
    X = np.log1p(X)
    return "text", balance(std(X), np.array(ys, dtype=np.int64), per=700, max_classes=2)


def prep_power():
    z = fetch("power")
    cand = [n for n in z.namelist() if "household_power" in n.lower() and n.lower().endswith(".txt")]
    raw = z.read(cand[0]).decode("latin-1")
    rows = [l.split(";") for l in raw.splitlines()[1:] if l.count(";") == 8 and "?" not in l]
    A = np.array([[float(v) for v in r[2:]] for r in rows])          # 7 个通道
    gap = A[:, 0]
    W = 60
    Xs, ys = [], []
    for i in range(0, len(A) - W - 1, 7):                            # 采样，别太密
        Xs.append(A[i:i + W].ravel())
        ys.append(gap[i + W])
    Xs = np.array(Xs); ys = np.array(ys)
    q = np.quantile(ys, [1 / 6, 2 / 6, 3 / 6, 4 / 6, 5 / 6])
    y = np.digitize(ys, q)
    return "timeseries", balance(std(Xs), y.astype(np.int64), per=700, max_classes=6)


def prep_ml100k():
    z = fetch("ml100k")
    u = np.array([[int(v) for v in l.decode("latin-1").split("\t")[:3]]
                  for l in read_zip(z, "u.data").splitlines() if l.strip()])
    items = []
    for l in read_zip(z, "u.item").splitlines():
        if not l.strip():
            continue
        p = l.decode("latin-1").split("|")
        items.append([float(v) for v in p[5:24]])
    items = np.array(items)
    nu = int(u[:, 0].max())
    uf = np.zeros((nu + 1, items.shape[1]))
    cnt = np.zeros((nu + 1, 1))
    for uu, ii, r in u[:, :3]:
        uf[uu] += items[ii - 1]; cnt[uu] += 1     # u.item 是 1-based
    uf = uf / np.maximum(cnt, 1)
    X = np.hstack([uf[u[:, 0]], items[u[:, 1] - 1]])
    r = u[:, 2]
    y = np.digitize(r, [3, 4, 5])          # 1-2 / 3 / 4 / 5 -> 4 类
    return "recommendation", balance(std(X), y.astype(np.int64), per=900, max_classes=4)


def prep_credit():
    z = fetch("credit")
    cand = [n for n in z.namelist() if n.lower().endswith(".xls") or n.lower().endswith(".xlsx")]
    raw = z.read(cand[0])
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    A = np.array([[float(v) for v in r[1:]] for r in rows[2:] if r[0] is not None])
    y = A[:, -1].astype(np.int64); X = A[:, :-1]
    return "tabular", balance(std(X), y, per=2000, max_classes=2)


def prep_cancer():
    z = fetch("cancer")
    cand = [n for n in z.namelist() if n.lower().endswith("wdbc.data")]
    raw = z.read(cand[0]).decode("latin-1")
    rows = [l.split(",") for l in raw.splitlines() if l.strip()]
    X = np.array([[float(v) for v in r[2:]] for r in rows])
    y = np.array([0 if r[1] == "B" else 1 for r in rows], dtype=np.int64)
    return "tabular", balance(std(X), y, per=None, max_classes=2)


PREPS = {"adult": prep_adult, "sms": prep_sms, "power": prep_power,
         "ml100k": prep_ml100k, "credit": prep_credit, "cancer": prep_cancer}

if __name__ == "__main__":
    only = sys.argv[1:] or list(PREPS)
    for name in only:
        try:
            nature, (X, y) = PREPS[name]()
            np.save(os.path.join(DATA, "ds_%s_X.npy" % name), X)
            np.save(os.path.join(DATA, "ds_%s_y.npy" % name), y)
            RESULTS[name] = dict(nature=nature, n=int(len(X)), dim=int(X.shape[1]),
                                 classes=int(len(set(y.tolist()))),
                                 dist=[int(v) for v in np.bincount(y)])
            print("OK   %-8s %-16s n=%-6d dim=%-5d classes=%d" % (name, nature, len(X), X.shape[1], len(set(y.tolist()))))
        except Exception as e:
            RESULTS[name] = dict(error=type(e).__name__ + ": " + str(e)[:160])
            print("FAIL %-8s %s: %s" % (name, type(e).__name__, str(e)[:160]))
        sys.stdout.flush()
    mf = os.path.join(DATA, "ds_manifest.json")
    old = json.load(open(mf, encoding="utf-8")) if os.path.exists(mf) else {}
    old.update(RESULTS)
    json.dump(old, open(mf, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("\n写好了 -> data/ds_*.npy 与 data/ds_manifest.json")
